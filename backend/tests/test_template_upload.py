import io

from docx import Document

from app.schemas.document_template import (
    TemplateClassificationResult,
    TemplateSectionClassification,
)
from app.services import template_upload_service


def _sample_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Purpose", level=1)
    doc.add_paragraph("This document captures meeting outcomes.")
    doc.add_heading("Requirements", level=1)
    doc.add_heading("Budget Notes", level=1)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _mock_classify(monkeypatch, mapping: dict[int, dict]):
    """mapping: section_index -> {content_type, custom_description}"""
    def fake_classify(self, sections):
        return TemplateClassificationResult(
            sections=[
                TemplateSectionClassification(
                    section_index=s.section_index,
                    content_type=mapping[s.section_index]["content_type"],
                    custom_description=mapping[s.section_index].get("custom_description"),
                    reasoning="mocked",
                )
                for s in sections
            ]
        )
    monkeypatch.setattr(
        template_upload_service.TemplateSchemaAgent, "classify", fake_classify
    )


def test_upload_template_classifies_sections(client, monkeypatch):
    _mock_classify(monkeypatch, {
        0: {"content_type": "custom_text"},
        1: {"content_type": "requirement"},
        2: {"content_type": "dynamic_custom", "custom_description": "Budget figures raised in the meeting."},
    })

    res = client.post(
        "/api/templates/upload",
        data={"name": "Test Template"},
        files={"file": ("test.docx", _sample_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["name"] == "Test Template"
    assert len(body["sections"]) == 3
    assert body["sections"][0]["content_type"] == "custom_text"
    assert body["sections"][0]["static_text"] == "This document captures meeting outcomes."
    assert body["sections"][1]["content_type"] == "requirement"
    assert body["sections"][2]["content_type"] == "dynamic_custom"
    assert body["sections"][2]["custom_description"] == "Budget figures raised in the meeting."


def test_upload_rejects_non_docx(client):
    res = client.post(
        "/api/templates/upload",
        data={"name": "Bad file"},
        files={"file": ("test.txt", b"not a docx", "text/plain")},
    )
    assert res.status_code == 400


def test_upload_rejects_template_with_no_headings(client, monkeypatch):
    doc = Document()
    doc.add_paragraph("Just plain text, no headings.")
    buf = io.BytesIO()
    doc.save(buf)

    res = client.post(
        "/api/templates/upload",
        data={"name": "No headings"},
        files={"file": ("test.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert res.status_code == 400


def test_dynamic_custom_without_description_falls_back_to_custom_text(client, monkeypatch):
    """Safety net: if the agent violates its own rule and omits
    custom_description for a dynamic_custom section, we must not run
    live extraction with empty criteria."""
    _mock_classify(monkeypatch, {
        0: {"content_type": "custom_text"},
        1: {"content_type": "requirement"},
        2: {"content_type": "dynamic_custom"},  # missing custom_description
    })

    res = client.post(
        "/api/templates/upload",
        data={"name": "Bad classification"},
        files={"file": ("test.docx", _sample_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert res.status_code == 200
    assert res.json()["sections"][2]["content_type"] == "custom_text"


def test_list_and_get_template(client, monkeypatch):
    _mock_classify(monkeypatch, {
        0: {"content_type": "custom_text"},
        1: {"content_type": "requirement"},
        2: {"content_type": "custom_text"},
    })
    upload_res = client.post(
        "/api/templates/upload",
        data={"name": "Listable"},
        files={"file": ("test.docx", _sample_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    template_id = upload_res.json()["id"]

    list_res = client.get("/api/templates")
    assert any(t["id"] == template_id for t in list_res.json())

    get_res = client.get(f"/api/templates/{template_id}")
    assert get_res.status_code == 200
    assert len(get_res.json()["sections"]) == 3


def test_get_unknown_template_returns_404(client):
    res = client.get("/api/templates/00000000-0000-0000-0000-000000000000")
    assert res.status_code == 404