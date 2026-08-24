from pathlib import Path

from docx import Document

from app.schemas.mom import (
    MOMDecision,
    MOMDocument,
    MOMRequirement,
)
from app.services.mom_document_service import (
    MOMDocumentService,
)


def create_test_mom():

    return MOMDocument(
        meeting_title="Authentication Meeting",
        meeting_date="2026-08-13",
        participants=[
            "Sarah",
            "John",
        ],
        requirements=[
            MOMRequirement(
                text=(
                    "The application must support "
                    "OTP-based login for all external users."
                )
            )
        ],
        decisions=[
            MOMDecision(
                text=(
                    "OTP authentication will be mandatory."
                )
            )
        ],
    )


def test_generate_mom_docx(tmp_path):

    mom = create_test_mom()

    output_path = (
        tmp_path / "Meeting_Minutes.docx"
    )

    service = MOMDocumentService()

    result = service.generate(
        mom=mom,
        output_path=str(output_path),
    )

    assert result == str(output_path)

    assert output_path.exists()

    assert output_path.is_file()


def test_generated_docx_can_be_opened(tmp_path):

    mom = create_test_mom()

    output_path = (
        tmp_path / "Meeting_Minutes.docx"
    )

    service = MOMDocumentService()

    service.generate(
        mom=mom,
        output_path=str(output_path),
    )

    document = Document(
        str(output_path)
    )

    full_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "Authentication Meeting" in full_text

    assert "Requirements" in full_text

    assert "Decisions" in full_text

    assert (
        "The application must support "
        "OTP-based login for all external users."
        in full_text
    )

    assert (
        "OTP authentication will be mandatory."
        in full_text
    )


def test_evidence_is_not_written_to_docx(tmp_path):

    mom = create_test_mom()

    output_path = (
        tmp_path / "Meeting_Minutes.docx"
    )

    service = MOMDocumentService()

    service.generate(
        mom=mom,
        output_path=str(output_path),
    )

    document = Document(
        str(output_path)
    )

    full_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    # Evidence-related fields must not appear
    # in the generated MOM.

    assert "SEG-0001" not in full_text

    assert "00:10:15" not in full_text

    assert "Evidence" not in full_text