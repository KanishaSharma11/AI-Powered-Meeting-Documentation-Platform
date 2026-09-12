import io
import uuid

from docx import Document

from app.models.document_template import DocumentTemplate, TemplateSection
from app.models.extraction import ExtractedItem
from app.models.fact import ProjectFact
from app.models.meeting import Meeting
from app.models.project import Project
from app.models.transcript import TranscriptSegment
from app.schemas.dynamic_extraction import DynamicSectionExtractionResult, DynamicExtractedItem
from app.schemas.extraction import EvidenceReference


def _standalone_meeting_with_items(db):
    meeting = Meeting(title="Standalone Meeting")
    db.add(meeting)
    db.flush()

    seg = TranscriptSegment(
        meeting_id=meeting.id, segment_id="SEG-0001", speaker="Sarah",
        timestamp="00:00:05", text="OTP login is required.",
        source_line_start=1, source_line_end=1,
    )
    db.add(seg)
    db.flush()

    item = ExtractedItem(
        meeting_id=meeting.id, item_type="requirement", agent_name="requirements_decisions_agent",
        text="The system shall support OTP login.", segment_id=seg.segment_id,
        evidence_speaker=seg.speaker, evidence_timestamp=seg.timestamp, evidence_text=seg.text,
        extraction_confidence=0.9, source_baseline_confidence=0.99, combined_confidence=0.891,
        review_status="accepted",
    )
    db.add(item)
    db.commit()
    return meeting


def _template_with_section(db, content_type, is_table=False, table_columns=None, custom_description=None, heading="Requirements"):
    template = DocumentTemplate(name="T", original_filename="t.docx")
    db.add(template)
    db.flush()
    db.add(TemplateSection(
        template_id=template.id, order_index=0, heading_text=heading, heading_level=1,
        content_type=content_type, is_table=is_table, table_columns=table_columns,
        custom_description=custom_description,
    ))
    db.commit()
    db.refresh(template)
    return template


def test_download_document_for_standalone_meeting(client, db):
    meeting = _standalone_meeting_with_items(db)
    template = _template_with_section(db, "requirement")

    res = client.get(f"/api/meetings/{meeting.id}/documents/{template.id}/download")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    doc = Document(io.BytesIO(res.content))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "OTP login" in full_text


def test_download_document_missing_meeting_returns_404(client, db):
    template = _template_with_section(db, "requirement")
    fake_id = "00000000-0000-0000-0000-000000000000"
    res = client.get(f"/api/meetings/{fake_id}/documents/{template.id}/download")
    assert res.status_code == 404


def test_download_document_missing_template_returns_404(client, db):
    meeting = _standalone_meeting_with_items(db)
    fake_id = "00000000-0000-0000-0000-000000000000"
    res = client.get(f"/api/meetings/{meeting.id}/documents/{fake_id}/download")
    assert res.status_code == 404


def test_table_section_maps_columns(client, db):
    meeting = Meeting(title="Action Meeting")
    db.add(meeting)
    db.flush()
    seg = TranscriptSegment(meeting_id=meeting.id, segment_id="SEG-0001", speaker="Raj",
        timestamp="00:00:05", text="Send contract by Friday.", source_line_start=1, source_line_end=1)
    db.add(seg)
    db.flush()
    item = ExtractedItem(
        meeting_id=meeting.id, item_type="action_item", agent_name="action_items_agent",
        text="Send the API contract.", owner="Raj", due_date="Friday", segment_id=seg.segment_id,
        evidence_speaker=seg.speaker, evidence_timestamp=seg.timestamp, evidence_text=seg.text,
        extraction_confidence=0.9, source_baseline_confidence=0.99, combined_confidence=0.891,
        review_status="accepted",
    )
    db.add(item)
    db.commit()

    template = _template_with_section(
        db, "action_item", is_table=True, table_columns=["Task", "Owner", "Due Date"], heading="Action Items"
    )

    res = client.get(f"/api/meetings/{meeting.id}/documents/{template.id}/download")
    assert res.status_code == 200
    doc = Document(io.BytesIO(res.content))
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert table.rows[1].cells[0].text == "Send the API contract."
    assert table.rows[1].cells[1].text == "Raj"
    assert table.rows[1].cells[2].text == "Friday"


def test_project_linked_meeting_pulls_from_project_facts(client, db):
    project = Project(name="Proj")
    db.add(project)
    db.flush()

    meeting = Meeting(title="Project Meeting", project_id=project.id)
    db.add(meeting)
    db.flush()

    fact = ProjectFact(
        project_id=project.id, fact_type="requirement", canonical_text="OTP login is mandatory.",
        status="active", confidence=0.9, first_meeting_id=meeting.id, last_meeting_id=meeting.id,
    )
    db.add(fact)
    db.commit()

    template = _template_with_section(db, "requirement")

    res = client.get(f"/api/meetings/{meeting.id}/documents/{template.id}/download")
    assert res.status_code == 200
    doc = Document(io.BytesIO(res.content))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "OTP login is mandatory." in full_text


def test_conflicted_project_fact_excluded_from_generated_document(client, db):
    project = Project(name="Proj2")
    db.add(project)
    db.flush()
    meeting = Meeting(title="Meeting2", project_id=project.id)
    db.add(meeting)
    db.flush()
    fact = ProjectFact(
        project_id=project.id, fact_type="requirement", canonical_text="Conflicted requirement.",
        status="conflicted", confidence=0.9, first_meeting_id=meeting.id, last_meeting_id=meeting.id,
    )
    db.add(fact)
    db.commit()

    template = _template_with_section(db, "requirement")
    res = client.get(f"/api/meetings/{meeting.id}/documents/{template.id}/download")
    doc = Document(io.BytesIO(res.content))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Conflicted requirement." not in full_text
    assert "No requirement recorded" in full_text


def test_dynamic_custom_section_runs_live_extraction(client, db, monkeypatch):
    meeting = Meeting(title="Budget Meeting")
    db.add(meeting)
    db.flush()
    seg = TranscriptSegment(meeting_id=meeting.id, segment_id="SEG-0001", speaker="Raj",
        timestamp="00:00:05", text="Vendor quoted 50000 for SMS gateway.", source_line_start=1, source_line_end=1)
    db.add(seg)
    db.commit()

    template = _template_with_section(
        db, "dynamic_custom", custom_description="Budget figures raised in the meeting.", heading="Budget Notes"
    )

    class FakeAgent:
        def __init__(self, criteria):
            self.criteria = criteria
        def extract(self, transcript):
            return DynamicSectionExtractionResult(items=[
                DynamicExtractedItem(text="Vendor quoted $50,000 for SMS gateway.",
                    evidence=EvidenceReference(segment_id="SEG-0001"), confidence=0.9)
            ])

    monkeypatch.setattr(
        "app.ai.agents.dynamic_section_agent.DynamicSectionAgent", FakeAgent
    )

    res = client.get(f"/api/meetings/{meeting.id}/documents/{template.id}/download")
    assert res.status_code == 200
    doc = Document(io.BytesIO(res.content))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "$50,000" in full_text
    assert "not yet human-reviewed" in full_text


def test_dynamic_custom_with_no_matching_content_is_honest(client, db, monkeypatch):
    meeting = Meeting(title="Unrelated Meeting")
    db.add(meeting)
    db.flush()
    seg = TranscriptSegment(meeting_id=meeting.id, segment_id="SEG-0001", speaker="Sarah",
        timestamp="00:00:05", text="Lets talk about login only.", source_line_start=1, source_line_end=1)
    db.add(seg)
    db.commit()

    template = _template_with_section(
        db, "dynamic_custom", custom_description="Budget figures raised in the meeting.", heading="Budget Notes"
    )

    class FakeAgentEmpty:
        def __init__(self, criteria):
            pass
        def extract(self, transcript):
            return DynamicSectionExtractionResult(items=[])

    monkeypatch.setattr(
        "app.ai.agents.dynamic_section_agent.DynamicSectionAgent", FakeAgentEmpty
    )

    res = client.get(f"/api/meetings/{meeting.id}/documents/{template.id}/download")
    assert res.status_code == 200
    doc = Document(io.BytesIO(res.content))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Nothing matching" in full_text


def test_custom_text_section_reproduces_static_text_verbatim(client, db):
    meeting = _standalone_meeting_with_items(db)
    template = DocumentTemplate(name="T2", original_filename="t2.docx")
    db.add(template)
    db.flush()
    db.add(TemplateSection(
        template_id=template.id, order_index=0, heading_text="Purpose", heading_level=1,
        content_type="custom_text", is_table=False,
        static_text="This document is confidential and for internal use only.",
    ))
    db.commit()
    db.refresh(template)

    res = client.get(f"/api/meetings/{meeting.id}/documents/{template.id}/download")
    doc = Document(io.BytesIO(res.content))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "confidential and for internal use only" in full_text