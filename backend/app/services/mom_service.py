"""
Document Generation Engine — Stage 1/2 fixed-format MOM (Section 2.2
step 6). Template-driven generation (Doc Intelligence Layer, Section
2.6) is a Stage 4 concern; for now this produces a clean, fixed MOM
layout from ACCEPTED/EDITED review items only.
"""

import io
from datetime import date

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.models.extraction import ExtractedItem
from app.models.meeting import Meeting
from app.models.transcript import TranscriptSegment
from app.schemas.mom import (
    MOMDocument,
    MOMRequirement,
    MOMDecision,
    MOMActionItem,
    MOMRiskOrQuestion,
)

_HEADING_COLOR = RGBColor(0x1F, 0x3B, 0x57)


def build_mom_document_model(
    meeting: Meeting, items: list[ExtractedItem], segments: list[TranscriptSegment]
) -> MOMDocument:
    """
    Build the clean MOMDocument model from only the items a human
    has approved (accepted or edited) — rejected and still-pending
    items never make it into the delivered document (Section 2.2
    step 7: "Human Review & Approval").
    """
    approved = [i for i in items if i.review_status in ("accepted", "edited")]

    participants = sorted({s.speaker for s in segments if s.speaker and s.speaker != "Unknown"})

    return MOMDocument(
        meeting_title=meeting.title,
        meeting_date=meeting.meeting_date.isoformat() if meeting.meeting_date else None,
        participants=participants,
        requirements=[
            MOMRequirement(text=i.final_text)
            for i in approved
            if i.item_type == "requirement"
        ],
        decisions=[
            MOMDecision(text=i.final_text)
            for i in approved
            if i.item_type == "decision"
        ],
        action_items=[
            MOMActionItem(text=i.final_text, owner=i.owner, due_date=i.due_date)
            for i in approved
            if i.item_type == "action_item"
        ],
        risks=[
            MOMRiskOrQuestion(text=i.final_text, category=i.item_type)
            for i in approved
            if i.item_type in ("risk", "blocker", "open_question")
        ],
    )


def _add_heading(doc: Document, text: str):
    h = doc.add_heading(text, level=1)
    for run in h.runs:
        run.font.color.rgb = _HEADING_COLOR


def render_mom_docx(mom: MOMDocument) -> bytes:
    doc = Document()

    title = doc.add_heading(mom.meeting_title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        f"Date: {mom.meeting_date or 'Not specified'}    |    Generated: {date.today().isoformat()}"
    )
    meta_run.italic = True
    meta_run.font.size = Pt(10)

    if mom.participants:
        _add_heading(doc, "Participants")
        doc.add_paragraph(", ".join(mom.participants))

    _add_heading(doc, "Requirements")
    if mom.requirements:
        for r in mom.requirements:
            doc.add_paragraph(r.text, style="List Bullet")
    else:
        doc.add_paragraph("No requirements recorded.")

    _add_heading(doc, "Decisions")
    if mom.decisions:
        for d in mom.decisions:
            doc.add_paragraph(d.text, style="List Bullet")
    else:
        doc.add_paragraph("No decisions recorded.")

    _add_heading(doc, "Action Items")
    if mom.action_items:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Action", "Owner", "Due Date"
        for a in mom.action_items:
            row = table.add_row().cells
            row[0].text = a.text
            row[1].text = a.owner or "Unassigned"
            row[2].text = a.due_date or "Not specified"
    else:
        doc.add_paragraph("No action items recorded.")

    _add_heading(doc, "Risks & Open Questions")
    if mom.risks:
        for r in mom.risks:
            label = r.category.replace("_", " ").title()
            doc.add_paragraph(f"[{label}] {r.text}", style="List Bullet")
    else:
        doc.add_paragraph("No risks or open questions recorded.")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()