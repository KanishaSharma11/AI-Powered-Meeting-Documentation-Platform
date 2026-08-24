from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt

from app.schemas.mom import MOMDocument


class MOMDocumentService:
    """
    Generates an editable Microsoft Word document from
    a clean MOMDocument.

    Evidence is intentionally NOT written into the document.
    Evidence remains in the application's internal intelligence
    layer and will later be accessible through the meeting chat.
    """

    def generate(
        self,
        mom: MOMDocument,
        output_path: str,
    ) -> str:
        """
        Generate a .docx MOM document.

        Parameters
        ----------
        mom:
            Clean MOMDocument containing meeting information,
            requirements, and decisions.

        output_path:
            Location where the Word document should be saved.

        Returns
        -------
        str
            Path to the generated Word document.
        """

        document = Document()

        self._configure_document(document)

        self._add_title(
            document,
            mom,
        )

        self._add_meeting_information(
            document,
            mom,
        )

        self._add_requirements(
            document,
            mom,
        )

        self._add_decisions(
            document,
            mom,
        )

        output = Path(output_path)

        output.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        document.save(output)

        return str(output)

    # ==========================================================
    # DOCUMENT CONFIGURATION
    # ==========================================================

    @staticmethod
    def _configure_document(
        document: Document,
    ) -> None:
        """
        Configure basic Word document properties.
        """

        section = document.sections[0]

        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

        styles = document.styles

        normal_style = styles["Normal"]

        normal_style.font.name = "Aptos"
        normal_style.font.size = Pt(11)

    # ==========================================================
    # TITLE
    # ==========================================================

    @staticmethod
    def _add_title(
        document: Document,
        mom: MOMDocument,
    ) -> None:
        """
        Add the main MOM title.
        """

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run(
            mom.meeting_title
        )

        run.bold = True
        run.font.size = Pt(20)

        subtitle = document.add_paragraph()

        subtitle.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = subtitle.add_run(
            "Minutes of Meeting"
        )

        run.italic = True
        run.font.size = Pt(11)

        document.add_paragraph()

    # ==========================================================
    # MEETING INFORMATION
    # ==========================================================

    @staticmethod
    def _add_meeting_information(
        document: Document,
        mom: MOMDocument,
    ) -> None:
        """
        Add meeting metadata.
        """

        heading = document.add_paragraph()

        run = heading.add_run(
            "Meeting Information"
        )

        run.bold = True
        run.font.size = Pt(14)

        table = document.add_table(
            rows=0,
            cols=2,
        )

        table.alignment = (
            WD_TABLE_ALIGNMENT.CENTER
        )

        table.style = "Table Grid"

        # Meeting title

        row = table.add_row().cells

        row[0].text = "Meeting"

        row[1].text = (
            mom.meeting_title
        )

        # Meeting date

        if mom.meeting_date:

            row = table.add_row().cells

            row[0].text = "Date"

            row[1].text = (
                mom.meeting_date
            )

        # Participants

        if mom.participants:

            row = table.add_row().cells

            row[0].text = "Participants"

            row[1].text = ", ".join(
                mom.participants
            )

        document.add_paragraph()

    # ==========================================================
    # REQUIREMENTS
    # ==========================================================

    @staticmethod
    def _add_requirements(
        document: Document,
        mom: MOMDocument,
    ) -> None:
        """
        Add requirements section.
        """

        heading = document.add_paragraph()

        run = heading.add_run(
            "1. Requirements"
        )

        run.bold = True
        run.font.size = Pt(14)

        if not mom.requirements:

            paragraph = document.add_paragraph()

            paragraph.add_run(
                "No requirements were identified."
            )

            document.add_paragraph()

            return

        for index, requirement in enumerate(
            mom.requirements,
            start=1,
        ):

            paragraph = document.add_paragraph(
                style="List Number"
            )

            run = paragraph.add_run(
                requirement.text
            )

            run.font.size = Pt(11)

        document.add_paragraph()

    # ==========================================================
    # DECISIONS
    # ==========================================================

    @staticmethod
    def _add_decisions(
        document: Document,
        mom: MOMDocument,
    ) -> None:
        """
        Add decisions section.
        """

        heading = document.add_paragraph()

        run = heading.add_run(
            "2. Decisions"
        )

        run.bold = True
        run.font.size = Pt(14)

        if not mom.decisions:

            paragraph = document.add_paragraph()

            paragraph.add_run(
                "No decisions were identified."
            )

            return

        for decision in mom.decisions:

            paragraph = document.add_paragraph(
                style="List Number"
            )

            run = paragraph.add_run(
                decision.text
            )

            run.font.size = Pt(11)