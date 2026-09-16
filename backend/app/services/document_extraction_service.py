"""
Extracts raw transcript text from an uploaded file, regardless of
format. Supported: .txt/.md (UTF-8, falls back to latin-1), .docx
(paragraphs + tables, in document order), .pdf (page text).
"""
import io
from docx import Document
from pypdf import PdfReader


class UnsupportedFileTypeError(ValueError):
    pass


class EmptyExtractionError(ValueError):
    pass


def _extract_txt(file_bytes: bytes) -> str:
    try:
        return file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return file_bytes.decode("latin-1")


def _extract_docx(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))
    lines: list[str] = []
    body = document.element.body
    paragraphs_by_id = {p._p: p for p in document.paragraphs}
    tables_by_id = {t._tbl: t for t in document.tables}
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p" and child in paragraphs_by_id:
            text = paragraphs_by_id[child].text.strip()
            if text:
                lines.append(text)
        elif tag == "tbl" and child in tables_by_id:
            for row in tables_by_id[child].rows:
                cell_texts = [c.text.strip() for c in row.cells if c.text.strip()]
                if len(cell_texts) >= 2:
                    # Two-column transcript tables are almost always
                    # "Speaker | Utterance" — join with a colon so the
                    # normalizer's speaker patterns actually match.
                    lines.append(f"{cell_texts[0]}: {' '.join(cell_texts[1:])}")
                elif cell_texts:
                    lines.append(cell_texts[0])
    return "\n".join(lines)


def _extract_pdf(file_bytes: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(f"Could not read PDF file: {e}") from e
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("This PDF is password-protected and cannot be read. Please upload an unprotected version.")
    pages_text = []
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if text:
            pages_text.append(text)
    return "\n\n".join(pages_text)


_EXTRACTORS = {".txt": _extract_txt, ".md": _extract_txt, ".docx": _extract_docx, ".pdf": _extract_pdf}


def extract_transcript_text(filename: str, file_bytes: bytes) -> str:
    lowered = filename.lower()
    ext = next((e for e in _EXTRACTORS if lowered.endswith(e)), None)
    if ext is None:
        raise UnsupportedFileTypeError(f"Unsupported file type for '{filename}'. Supported: {', '.join(sorted(_EXTRACTORS))}.")
    text = _EXTRACTORS[ext](file_bytes)
    if not text or not text.strip():
        raise EmptyExtractionError(
            f"No extractable text found in '{filename}'. If this is a scanned/image-only PDF, "
            "text extraction isn't supported — try an exported text transcript instead."
        )
    return text