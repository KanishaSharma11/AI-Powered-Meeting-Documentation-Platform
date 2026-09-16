"""
Dynamic Document Generator (Section 2.6/2.7). Renders a .docx that
mirrors an uploaded template's structure, populated with real content:
  - requirement/decision/action_item/risk_or_question: from a
    project's ACTIVE ProjectFacts (Stage 3) if the meeting belongs to
    a project, else this meeting's own accepted/edited ExtractedItems.
  - participants: from transcript segment speakers.
  - custom_text: template's own static_text reproduced verbatim.
  - dynamic_custom: live evidence-grounded extraction, rendered
    italicized with a provenance note since it skips the Stage 2
    human review gate.
"""
import io
from dataclasses import dataclass
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from sqlalchemy.orm import Session

from app.models.document_template import DocumentTemplate
from app.models.extraction import ExtractedItem
from app.models.fact import ProjectFact
from app.models.meeting import Meeting
from app.models.transcript import TranscriptSegment
from app.services.fusion_service import FUSION_TYPE_GROUPS
from app.services.normalization import render_segments_for_prompt

_HEADING_COLOR = RGBColor(0x1F, 0x3B, 0x57)

_CONTENT_TYPE_TO_ITEM_TYPES: dict[str, list[str]] = {}
for _item_type, _bucket in FUSION_TYPE_GROUPS.items():
    _CONTENT_TYPE_TO_ITEM_TYPES.setdefault(_bucket, []).append(_item_type)


@dataclass
class ContentEntry:
    text: str
    owner: str | None = None
    due_date: str | None = None


def _resolve_participants(db: Session, meeting: Meeting) -> list[str]:
    if meeting.project_id is not None:
        rows = (
            db.query(TranscriptSegment.speaker)
            .join(Meeting, Meeting.id == TranscriptSegment.meeting_id)
            .filter(Meeting.project_id == meeting.project_id)
            .distinct()
            .all()
        )
    else:
        rows = db.query(TranscriptSegment.speaker).filter(TranscriptSegment.meeting_id == meeting.id).distinct().all()
    return sorted({r[0] for r in rows if r[0] and r[0] != "Unknown"})


def _resolve_content(db: Session, meeting: Meeting, content_type: str) -> list[ContentEntry]:
    if meeting.project_id is not None:
        facts = (
            db.query(ProjectFact)
            .filter(ProjectFact.project_id == meeting.project_id, ProjectFact.fact_type == content_type, ProjectFact.status == "active")
            .order_by(ProjectFact.created_at)
            .all()
        )
        return [ContentEntry(text=f.canonical_text, owner=f.owner, due_date=f.due_date) for f in facts]

    item_types = _CONTENT_TYPE_TO_ITEM_TYPES.get(content_type, [])
    items = (
        db.query(ExtractedItem)
        .filter(ExtractedItem.meeting_id == meeting.id, ExtractedItem.item_type.in_(item_types), ExtractedItem.review_status.in_(("accepted", "edited")))
        .order_by(ExtractedItem.created_at)
        .all()
    )
    return [ContentEntry(text=i.final_text, owner=i.owner, due_date=i.due_date) for i in items]


def _match_column(column_name: str, entry: ContentEntry) -> str:
    lowered = column_name.lower()
    if "owner" in lowered or "assignee" in lowered or "responsible" in lowered:
        return entry.owner or "Unassigned"
    if "due" in lowered or "date" in lowered or "deadline" in lowered:
        return entry.due_date or "Not specified"
    return entry.text


def _resolve_dynamic_content(db: Session, meeting: Meeting, criteria: str) -> list[ContentEntry]:
    from app.ai.agents.dynamic_section_agent import DynamicSectionAgent

    segments = db.query(TranscriptSegment).filter(TranscriptSegment.meeting_id == meeting.id).order_by(TranscriptSegment.source_line_start).all()
    if not segments:
        return []

    transcript_block = render_segments_for_prompt(segments)
    result = DynamicSectionAgent(criteria=criteria).extract(transcript_block)

    segments_by_id = {s.segment_id: s for s in segments}
    entries = []
    for item in result.items:
        if item.evidence.segment_id in segments_by_id:
            entries.append(ContentEntry(text=item.text))
    return entries


def generate_document(db: Session, meeting: Meeting, template: DocumentTemplate) -> bytes:
    doc = Document()
    title = doc.add_heading(meeting.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for section in template.sections:
        level = max(1, min(section.heading_level, 9))
        heading = doc.add_heading(section.heading_text, level=level)
        for run in heading.runs:
            run.font.color.rgb = _HEADING_COLOR

        if section.content_type == "custom_text":
            text = section.static_text or ""
            if text.strip():
                for line in text.split("\n"):
                    if line.strip():
                        doc.add_paragraph(line.strip())
            else:
                doc.add_paragraph("Not specified in template.")
            continue

        if section.content_type == "participants":
            participants = _resolve_participants(db, meeting)
            doc.add_paragraph(", ".join(participants) if participants else "No participants recorded.")
            continue

        if section.content_type == "dynamic_custom":
            criteria = section.custom_description or section.heading_text
            entries = _resolve_dynamic_content(db, meeting, criteria)

            note = doc.add_paragraph()
            note_run = note.add_run("AI-drafted from the meeting transcript — not yet human-reviewed.")
            note_run.italic = True
            note_run.font.size = Pt(9)
            note_run.font.color.rgb = RGBColor(0x8A, 0x8A, 0x8A)

            if not entries:
                p = doc.add_paragraph()
                r = p.add_run(f"Nothing matching \u201c{criteria}\u201d was identified in this meeting.")
                r.italic = True
                continue

            for entry in entries:
                p = doc.add_paragraph(style="List Bullet")
                r = p.add_run(entry.text)
                r.italic = True
            continue

        entries = _resolve_content(db, meeting, section.content_type)

        if not entries:
            doc.add_paragraph(f"No {section.content_type.replace('_', ' ')} recorded.")
            continue

        if section.is_table and section.table_columns:
            table = doc.add_table(rows=1, cols=len(section.table_columns))
            table.style = "Light Grid Accent 1"
            for i, col_name in enumerate(section.table_columns):
                table.rows[0].cells[i].text = col_name
            for entry in entries:
                row = table.add_row().cells
                for i, col_name in enumerate(section.table_columns):
                    row[i].text = _match_column(col_name, entry)
        else:
            for entry in entries:
                doc.add_paragraph(entry.text, style="List Bullet")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()