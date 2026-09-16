"""
Template Intelligence Layer, structural half: parses an uploaded
.docx template into an ordered list of raw sections. Pure structural
parsing, no AI — classification is a separate step (template_schema_agent.py).
"""
import io
from dataclasses import dataclass
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


@dataclass
class RawTemplateSection:
    order_index: int
    heading_text: str
    heading_level: int
    is_table: bool = False
    table_columns: list[str] | None = None
    raw_text: str = ""


def _heading_level(paragraph: Paragraph) -> int | None:
    style_name = (paragraph.style.name or "") if paragraph.style else ""
    if style_name.lower().startswith("heading"):
        digits = "".join(ch for ch in style_name if ch.isdigit())
        return int(digits) if digits else 1
    if style_name.lower() == "title":
        return 1
    return None


def _iter_body_blocks(document: Document):
    body = document.element.body
    paragraphs_by_id = {p._p: p for p in document.paragraphs}
    tables_by_id = {t._tbl: t for t in document.tables}
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p" and child in paragraphs_by_id:
            yield paragraphs_by_id[child]
        elif tag == "tbl" and child in tables_by_id:
            yield tables_by_id[child]


def parse_docx_template(file_bytes: bytes) -> list[RawTemplateSection]:
    document = Document(io.BytesIO(file_bytes))
    sections: list[RawTemplateSection] = []
    current: RawTemplateSection | None = None
    order = 0
    for block in _iter_body_blocks(document):
        if isinstance(block, Paragraph):
            level = _heading_level(block)
            text = block.text.strip()
            if level is not None and text:
                current = RawTemplateSection(order_index=order, heading_text=text, heading_level=level)
                sections.append(current)
                order += 1
                continue
            if current is not None and text:
                current.raw_text = (current.raw_text + "\n" + text).strip()
        elif isinstance(block, Table):
            if current is not None and not current.is_table:
                header_cells = block.rows[0].cells if block.rows else []
                current.is_table = True
                current.table_columns = [c.text.strip() for c in header_cells]
    return sections